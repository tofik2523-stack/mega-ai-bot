import os
import io
import json
import asyncio
import logging
from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import CommandStart
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardMarkup, InlineKeyboardButton, BufferedInputFile
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
import google.generativeai as genai
from pptx import Presentation
from docx import Document

# Logging
logging.basicConfig(level=logging.INFO)

# --- KONFIGURATSIYA (Sizning kalitlaringiz) ---
BOT_TOKEN = "8986220465:AAFn2dhcavFcMR7jrpcr4K_Hjzx3dfgjUms"
GEMINI_API_KEY = "AQ.Ab8RN6JKdvXIdQEtvk4L-Yy9gp6PjC-PEz-wdW_atPmGofTCPg"

# Gemini sozlanmasi
genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel('gemini-1.5-flash')

# Bot & Dispatcher
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

# FSM (Holatlar)
class Form(StatesGroup):
    waiting_pptx_topic = State()
    waiting_docx_topic = State()
    waiting_quiz_topic = State()
    waiting_text_to_slide = State()

# --- MENYULAR (UI/UX) ---
main_keyboard = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="🎓 Akademik va Ta'lim"), KeyboardButton(text="🎨 Media va AI Visual")],
        [KeyboardButton(text="📁 Hujjat va PDF Master"), KeyboardButton(text="💼 Biznes va Utility")],
        [KeyboardButton(text="📚 Mening Kutubxonam"), KeyboardButton(text="⚙️ Sozlamalar va Balans")]
    ],
    resize_keyboard=True
)

academic_inline = InlineKeyboardMarkup(
    inline_keyboard=[
        [InlineKeyboardButton(text="📊 Prezentatsiya (PPTX)", callback_data="cmd_pptx"), InlineKeyboardButton(text="📝 Referat (DOCX)", callback_data="cmd_docx")],
        [InlineKeyboardButton(text="🎯 Quiz & Test", callback_data="cmd_quiz"), InlineKeyboardButton(text="📑 Matnni Slaydga Aylantirish", callback_data="cmd_text2pptx")],
        [InlineKeyboardButton(text="🏠 Bosh menyu", callback_data="cmd_home")]
    ]
)

# --- HANDLERS ---

@dp.message(CommandStart())
async def start_handler(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer(
        f"Salom, **{message.from_user.first_name}**! 👋\n\n"
        f"Men **Mega AI Bot**man — sizning akademik, visual va biznes yordamchingizman.\n"
        f"Kerakli bo'limni pastdagi menyudan tanlang:",
        reply_markup=main_keyboard,
        parse_mode="Markdown"
    )

@dp.message(F.text == "🎓 Akademik va Ta'lim")
async def academic_menu(message: types.Message):
    await message.answer(
        "🎓 **Akademik va Ta'lim bo'limi**\n\nQuyidagi xizmatlardan birini tanlang:",
        reply_markup=academic_inline,
        parse_mode="Markdown"
    )

@dp.callback_query(F.data == "cmd_home")
async def home_callback(call: types.CallbackQuery, state: FSMContext):
    await state.clear()
    await call.message.answer("🏠 Bosh menyuga qaytdingiz:", reply_markup=main_keyboard)
    await call.answer()

# --- 1. PREZENTATSIYA (PPTX) GENERATORI ---
@dp.callback_query(F.data == "cmd_pptx")
async def pptx_start(call: types.CallbackQuery, state: FSMContext):
    await state.set_state(Form.waiting_pptx_topic)
    await call.message.answer("📊 **Prezentatsiya yaratish uchun mavzu kiriting:**\n\n*Masalan: O'zbekistonda sun'iy intellekt kelajagi*", parse_mode="Markdown")
    await call.answer()

@dp.message(Form.waiting_pptx_topic)
async def pptx_process(message: types.Message, state: FSMContext):
    topic = message.text
    msg = await message.answer("⏳ **AI slaydlar va strukturani tayyorlamoqda...**")
    
    prompt = f"""
    Mavzu: '{topic}'. Ushbu mavzu bo'yicha 5 ta slayddan iborat prezentatsiya uchun JSON formatda ma'lumot tayyorla.
    Format faqat sof JSON bo'lsin:
    [
      {{"title": "Slayd 1 Sarlavha", "content": ["Nuqta 1", "Nuqta 2", "Nuqta 3"]}},
      {{"title": "Slayd 2 Sarlavha", "content": ["Nuqta 1", "Nuqta 2"]}}
    ]
    Javobda hech qanday qo'shimcha matn yoki ```json belgisi bo'lmasin!
    """
    
    try:
        response = model.generate_content(prompt)
        text_res = response.text.strip().replace("```json", "").replace("```", "")
        slides_data = json.loads(text_res)
        
        prs = Presentation()
        for slide_info in slides_data:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = slide_info.get("title", "Slayd")
            content_shape = slide.placeholders[1]
            tf = content_shape.text_frame
            tf.word_wrap = True
            
            for i, point in enumerate(slide_info.get("content", [])):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = point
                
        file_stream = io.BytesIO()
        prs.save(file_stream)
        file_stream.seek(0)
        
        input_file = BufferedInputFile(file_stream.read(), filename=f"{topic[:20]}.pptx")
        await msg.delete()
        await message.answer_document(input_file, caption=f"✅ **'{topic}'** mavzusidagi tayyor prezentatsiyangiz!", parse_mode="Markdown")
    except Exception as e:
        await msg.edit_text(f"❌ Xatolik yuz berdi: {e}")
    
    await state.clear()

# --- 2. REFERAT (DOCX) GENERATORI ---
@dp.callback_query(F.data == "cmd_docx")
async def docx_start(call: types.CallbackQuery, state: FSMContext):
    await state.set_state(Form.waiting_docx_topic)
    await call.message.answer("📝 **Referat/Mustaqil ish uchun mavzu kiriting:**\n\n*Masalan: Raqamli iqtisodiyot va uning ahamiyati*", parse_mode="Markdown")
    await call.answer()

@dp.message(Form.waiting_docx_topic)
async def docx_process(message: types.Message, state: FSMContext):
    topic = message.text
    msg = await message.answer("⏳ **AI referat matni va rejasini shakllantirmoqda...**")
    
    prompt = f"Mavzu: '{topic}'. Ushbu mavzuda to'liq, akademik va batafsil Referat yozib ber. Reja, Kirish, Asosiy qism (2 ta bob) va Xulosa bo'lsin."
    
    try:
        response = model.generate_content(prompt)
        doc = Document()
        doc.add_heading(topic.upper(), 0)
        doc.add_paragraph(response.text)
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        input_file = BufferedInputFile(file_stream.read(), filename=f"Referat_{topic[:15]}.docx")
        await msg.delete()
        await message.answer_document(input_file, caption=f"✅ **'{topic}'** bo'yicha tayyor Referat (Word)!", parse_mode="Markdown")
    except Exception as e:
        await msg.edit_text(f"❌ Xatolik yuz berdi: {e}")
        
    await state.clear()

# --- 3. QUIZ & TEST GENERATORI ---
@dp.callback_query(F.data == "cmd_quiz")
async def quiz_start(call: types.CallbackQuery, state: FSMContext):
    await state.set_state(Form.waiting_quiz_topic)
    await call.message.answer("🎯 **Qaysi mavzuda test tayyorlaylik?**\n\n*Masalan: Informatika fanidan 5-sinf testlari*", parse_mode="Markdown")
    await call.answer()

@dp.message(Form.waiting_quiz_topic)
async def quiz_process(message: types.Message, state: FSMContext):
    topic = message.text
    msg = await message.answer("⏳ **Test savollari tuzilmoqda...**")
    
    prompt = f"""
    Mavzu: '{topic}'. Ushbu mavzuda 3 ta ko'p tanlovli (Quiz) test savoli tuz.
    Javob sof JSON bo'lsin:
    [
      {{
        "question": "Savol matni?",
        "options": ["A variant", "B variant", "C variant", "D variant"],
        "correct_option_id": 0
      }}
    ]
    Javobda hech qanday ```json yoki ortiqcha belgi bo'lmasin!
    """
    
    try:
        response = model.generate_content(prompt)
        text_res = response.text.strip().replace("```json", "").replace("```", "")
        quiz_data = json.loads(text_res)
        
        await msg.delete()
        for q in quiz_data:
            await message.answer_poll(
                question=q["question"],
                options=q["options"],
                type="quiz",
                correct_option_id=q["correct_option_id"],
                is_anonymous=False
            )
    except Exception as e:
        await msg.edit_text(f"❌ Test tuzishda xatolik: {e}")
        
    await state.clear()

# --- 4. MATNNI SLAYDGA AYLANTIRISH ---
@dp.callback_query(F.data == "cmd_text2pptx")
async def text2pptx_start(call: types.CallbackQuery, state: FSMContext):
    await state.set_state(Form.waiting_text_to_slide)
    await call.message.answer("📑 **Slaydga aylantirish kerak bo'lgan uzun matningizni yuboring:**", parse_mode="Markdown")
    await call.answer()

@dp.message(Form.waiting_text_to_slide)
async def text2pptx_process(message: types.Message, state: FSMContext):
    user_text = message.text
    msg = await message.answer("⏳ **Matn tahlil qilinib, slaydga joylanmoqda...**")
    
    prompt = f"Ushbu matnni tahlil qilib, mantiqiy bo'lingan 3-5 ta slayd strukturasi (JSON) ko'rinishida ber:\n{user_text}\nJSON: [{{\"title\": \"...\", \"content\": [\"...\"]}}]"
    
    try:
        response = model.generate_content(prompt)
        text_res = response.text.strip().replace("```json", "").replace("```", "")
        slides_data = json.loads(text_res)
        
        prs = Presentation()
        for slide_info in slides_data:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = slide_info.get("title", "Slayd")
            content_shape = slide.placeholders[1]
            tf = content_shape.text_frame
            tf.word_wrap = True
            for i, point in enumerate(slide_info.get("content", [])):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = point
                
        file_stream = io.BytesIO()
        prs.save(file_stream)
        file_stream.seek(0)
        
        input_file = BufferedInputFile(file_stream.read(), filename="Matndan_Slayd.pptx")
        await msg.delete()
        await message.answer_document(input_file, caption="✅ Matningiz muvaffaqiyatli Prezentatsiyaga aylantirildi!", parse_mode="Markdown")
    except Exception as e:
        await msg.edit_text(f"❌ Xatolik: {e}")
        
    await state.clear()

# --- ISHGA TUSHIRISH ---
async def main():
    print("🚀 Mega AI Bot ishga tushdi!")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
